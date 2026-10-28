"""P10 の次に差し込む「先行施策｜プリザーブドフラワーの卸販売」ページを PPTX（1枚）で組む。
リボン・アスタリスクは元資料の該当ページ（300dpi PNG）から切り出し、文字は Yu Gothic のテキストボックス。"""
import sys
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

ORIG = sys.argv[1]   # 元PDF の任意ページ（リボン・アスタリスク付き）の 300dpi PNG
OUT = sys.argv[2]
PAGE_NO = sys.argv[3] if len(sys.argv) > 3 else "11"
SC = 300 / 72

BROWN = RGBColor(0x60, 0x38, 0x2C)
SALMON = RGBColor(0xB9, 0x5B, 0x45)
ACCENT = RGBColor(0xF0, 0x8A, 0x70)
GRAY = RGBColor(0x4B, 0x4B, 0x4B)
DARK = RGBColor(0x30, 0x34, 0x3B)
SUB = RGBColor(0x5F, 0x66, 0x70)
PINK = RGBColor(0xFB, 0xEA, 0xE7)
BORDER = RGBColor(0xE6, 0xE1, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Yu Gothic"


def crop_pt(im, x0, y0, x1, y1, out):
    """白地の元ページから切り出し、白を透過にする。"""
    c = im.crop((int(x0 * SC), int(y0 * SC), int(x1 * SC), int(y1 * SC))).convert("RGBA")
    px = c.load()
    for yy in range(c.height):
        for xx in range(c.width):
            r, g, b, _ = px[xx, yy]
            if r > 246 and g > 246 and b > 246:
                px[xx, yy] = (255, 255, 255, 0)
    c.save(out)


def rounded(src, w_pt, h_pt, r_pt, out, px=4):
    im = Image.open(src).convert("RGB")
    W, H = int(w_pt * px), int(h_pt * px)
    sw, sh = im.size
    s = max(W / sw, H / sh)
    im = im.resize((max(W, int(sw * s + 0.5)), max(H, int(sh * s + 0.5))), Image.LANCZOS)
    l, t = (im.width - W) // 2, (im.height - H) // 2
    im = im.crop((l, t, l + W, t + H)).convert("RGBA")
    m = Image.new("L", (W, H), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, W - 1, H - 1), radius=int(r_pt * px), fill=255)
    im.putalpha(m)
    im.save(out)


orig = Image.open(ORIG).convert("RGB")
crop_pt(orig, 0, 40, 42, 510, "_ribbon.png")
crop_pt(orig, 866, 18, 936, 86, "_asterisk.png")

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(959.76), Pt(540)
s = prs.slides.add_slide(prs.slide_layouts[6])


def pic(path, x, y, w, h):
    return s.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))


def _font(run, size, color, bold=False, name=FONT, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, name
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)
    if spacing is not None:
        rpr.set("spc", str(int(spacing * 100)))


def text(x, y, w, h, runs, wrap=True, line=1.0, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    for t, size, color, bold, sp in runs:
        for i, part in enumerate(t.split("\n")):
            if i > 0:
                p = tf.add_paragraph()
                p.alignment = align
                p.line_spacing = line
            r = p.add_run()
            r.text = part
            _font(r, size, color, bold, spacing=sp)
    return tb


def shape(kind, x, y, w, h, fill=None, line=None, radius=None):
    r = s.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
    if radius is not None:
        r.adjustments[0] = radius / min(w, h)
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid(); r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line; r.line.width = Pt(0.75)
    r.shadow.inherit = False
    st = r._element.find(qn("p:style"))
    if st is not None:
        r._element.remove(st)
    return r


# ---- 固定要素 ----
pic("_ribbon.png", 0, 40, 42, 470)
pic("_asterisk.png", 866, 18, 70, 68)
text(56, 30, 760, 32, [("先行施策｜プリザーブドフラワーの卸販売", 23.8, BROWN, False, 1.5)], wrap=False)
tb = text(880, 506, 40, 20, [(PAGE_NO, 13.7, GRAY, False, 0)], wrap=False, align=PP_ALIGN.RIGHT)
tb.text_frame.paragraphs[0].runs[0].font.name = "Arial"

# ---- リード ----
text(56, 91, 870, 22, [("クリスマス・バレンタインから小さく始める先行施策です。Domuzが商品を卸し、貴社倉庫から発送いただきます。",
                        15.5, SALMON, True, 0)], wrap=False)

# ---- 左：商品候補 ----
text(56, 122, 480, 16, [("商品候補（プリザーブド・ドライ／価格は下代・税抜）", 13.7, SUB, True, 0)], wrap=False)
products = [
    ("glass_single.jpg", "ガラスボックス（一輪）", "プリザーブド\nM 直径4×高さ16cm\nL 直径5×高さ21cm", "M 950円\nL 1,200円"),
    ("plastic_box.jpg", "プラスチックボックス", "プリザーブド・ドライ\n直径5×高さ18cm", "1,300円"),
    ("dome_box.jpg", "ドームボックス", "プリザーブド・ドライ\n直径9×高さ7.7cm", "1,500円"),
]
cw, ch, gap, top = 152, 246, 12, 142
for i, (img, name, spec, price) in enumerate(products):
    x = 56 + i * (cw + gap)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, cw, ch, fill=WHITE, line=BORDER, radius=6)
    rounded(img, cw - 20, 128, 4, f"_p{i}.png")
    pic(f"_p{i}.png", x + 10, top + 10, cw - 20, 128)
    text(x + 10, top + 148, cw - 20, 18, [(name, 12, DARK, True, 0)], wrap=False)
    text(x + 10, top + 168, cw - 20, 34, [(spec, 9, SUB, False, 0)], line=1.25)
    text(x + 10, top + 206, cw - 20, 34, [(price, 13.5, BROWN, True, 0)], wrap=False, line=1.15)
text(56, 396, 480, 14, [("※ このほかガラスボックス（ブーケ）1,600円、スクエアボックス 1,900円などもご用意しています。", 9.5, SUB, False, 0)], wrap=False)

# ---- 右：進め方 ----
RX, RW = 566, 352
text(RX, 122, RW, 16, [("進め方（物の流れ）", 13.7, SUB, True, 0)], wrap=False)
flow = [("Domuz", "商品を製作し\n貴社倉庫へ納品"), ("貴社倉庫", "保管し、注文ごとに\n発送"), ("受け手", "プリザーブドの\nギフトが届く")]
bw, bh, ar = 104, 56, 20
for i, (t, sub) in enumerate(flow):
    x = RX + i * (bw + ar)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, 142, bw, bh, fill=PINK, radius=5)
    text(x + 8, 146, bw - 16, 16, [(t, 11.5, DARK, True, 0)], wrap=False)
    text(x + 8, 163, bw - 16, 32, [(sub, 8.5, SUB, False, 0)], line=1.2)
    if i < 2:
        shape(MSO_SHAPE.CHEVRON, x + bw + 3, 142 + 20, 14, 16, fill=ACCENT)

box_top, box_h = 212, 176
shape(MSO_SHAPE.ROUNDED_RECTANGLE, RX, box_top, RW, box_h, fill=WHITE, line=BORDER, radius=6)
rows = [
    ("想定シーン", "クリスマス、バレンタイン・ホワイトデー"),
    ("発送", "貴社倉庫から発送（Domuzは納品まで）"),
    ("情報連携", "不要。受注・有効化システムとの接続はありません"),
    ("取引条件", "最低ロット100個〜、請求書払い、送料別"),
    ("品質", "保存状態が良ければ1〜3年美しさを保ちます"),
]
for i, (k, v) in enumerate(rows):
    y = box_top + 12 + i * 32
    text(RX + 14, y + 1, 62, 16, [(k, 10.5, BROWN, True, 0)], wrap=False)
    bold = (k == "発送")
    text(RX + 80, y, RW - 94, 30, [(v, 10.5, DARK, bold, 0)], line=1.2)

# ---- 締めの一文 ----
text(56, 420, 862, 22, [("発送は貴社倉庫から行っていただきます。Domuzは商品の製作と貴社倉庫への一括納品までを担当します。",
                         15.9, SALMON, True, 0)], wrap=False, align=PP_ALIGN.CENTER)
# ---- 脚注 ----
text(56, 486, 862, 16, [("※ 価格はすべて下代（卸値・税抜）、送料別。販売価格（上代）は貴社で自由に設定いただけます。箔押し等のカスタマイズは別途ご相談ください。",
                         10.8, BROWN, False, 0.3)], wrap=False)

prs.save(OUT)
print("saved", OUT)

"""4ページ目（株式会社Domuzについて）を編集可能な PPTX として組む。
元PDFのリボン・アスタリスクは元ページから切り出した画像、文字は Yu Gothic のテキストボックス。"""
import sys
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

ORIG4 = sys.argv[1]  # 元PDF 4ページ目の 300dpi レンダリング
OUT = sys.argv[2]
SC = 300 / 72

BROWN = RGBColor(0x60, 0x38, 0x2C)
SALMON = RGBColor(0xB9, 0x5B, 0x45)
GRAY = RGBColor(0x4B, 0x4B, 0x4B)
DARK = RGBColor(0x30, 0x34, 0x3B)
SUB = RGBColor(0x5F, 0x66, 0x70)
PINK = RGBColor(0xFB, 0xEA, 0xE7)
FONT = "Yu Gothic"


def crop_pt(im, x0, y0, x1, y1, out):
    im.crop((int(x0 * SC), int(y0 * SC), int(x1 * SC), int(y1 * SC))).save(out)


def rounded(src, w_pt, h_pt, r_pt, out, px=4):
    im = Image.open(src).convert("RGB")
    W, H = int(w_pt * px), int(h_pt * px)
    sw, sh = im.size
    s = max(W / sw, H / sh)
    im = im.resize((max(W, int(sw * s + 0.5)), max(H, int(sh * s + 0.5))), Image.LANCZOS)
    l, t = (im.width - W) // 2, (im.height - H) // 2
    im = im.crop((l, t, l + W, t + H)).convert("RGBA")
    m = Image.new("L", (W, H), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, W - 1, H - 1), radius=int(r_pt * px), fill=255)
    im.putalpha(m)
    im.save(out)


orig = Image.open(ORIG4).convert("RGB")
crop_pt(orig, 0, 40, 42, 510, "_ribbon.png")
crop_pt(orig, 866, 18, 936, 86, "_asterisk.png")
for i, f in enumerate(["bq4.jpg", "bq2.jpg", "bq3.jpg"]):
    rounded(f, 104, 176, 6, f"_photo{i}.png")
for i, f in enumerate(["andflower.jpg", "wh1.jpg", "wh2.jpg"]):
    rounded(f, 78, 78, 4, f"_card{i}.png")

prs = Presentation()
prs.slide_width, prs.slide_height = Pt(959.76), Pt(540)
s = prs.slides.add_slide(prs.slide_layouts[6])


def pic(path, x, y, w, h):
    return s.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))


def _font(run, size, color, bold=False, name=FONT, spacing=None):
    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, name
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)
    if spacing is not None:
        rpr.set("spc", str(int(spacing * 100)))


def text(x, y, w, h, runs, wrap=True, line=1.0, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """runs: [(text, size, color, bold, spacing), ...] / 改行は '\n' を含む run で。"""
    tb = s.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line
    for t, size, color, bold, sp in runs:
        parts = t.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                p = tf.add_paragraph()
                p.alignment = align
                p.line_spacing = line
            r = p.add_run()
            r.text = part
            _font(r, size, color, bold, spacing=sp)
    return tb


# ---- 元ページの固定要素 ----
pic("_ribbon.png", 0, 40, 42, 470)
pic("_asterisk.png", 866, 18, 70, 68)
text(56, 30, 700, 32, [("株式会社Domuzについて", 23.8, BROWN, True, 1.5)], wrap=False)
text(880, 506, 40, 20, [("4", 13.7, GRAY, False, 0)], wrap=False, align=PP_ALIGN.RIGHT)
s.shapes[-1].text_frame.paragraphs[0].runs[0].font.name = "Arial"

# ---- リード ----
text(56, 91, 870, 22, [("お花のD2C「AND FLOWER」を運営し、川崎の自社拠点で花の仕入れからブーケの製作、発送までを一貫して手がけています。",
                        15, SALMON, True, 0)], wrap=False)

# ---- 会社概要（P2 と同じ書式） ----
text(56, 134, 400, 32, [("株式会社 ", 20.2, BROWN, True, 1), ("Domuz", 23.8, BROWN, True, 1)], wrap=False)
rows = [("代表者", "髙木 弘貴", None), ("設立", "2018年10月11日", None), ("従業員数", "69名", "2026年8月時点"),
        ("本社", "神奈川県川崎市", None), ("事業", "花と植物のD2C、花卉産業マーケットプレイス", None)]
for i, (k, v, note) in enumerate(rows):
    y = 178 + i * 37
    text(56, y + 4, 84, 22, [(k, 15.9, BROWN, True, 1)], wrap=False)
    runs = [(v, 20.2, GRAY, False, 1)]
    if note:
        runs.append(("　" + note, 7.9, GRAY, False, 1.5))
    text(140, y, 450, 28, runs, wrap=False, anchor=MSO_ANCHOR.BOTTOM)

# ---- 写真 ----
for i in range(3):
    pic(f"_photo{i}.png", 590 + i * 112, 134, 104, 176)
text(590, 316, 328, 16, [("【AND FLOWER ブーケ・アレンジメント製作例】", 10.8, BROWN, False, 0.5)],
     wrap=False, align=PP_ALIGN.CENTER)

# ---- カード ----
cards = [("お花のD2C「AND FLOWER」", "2021年より運営。ギフト向けのブーケ・アレンジメントをオンラインで届けています。"),
         ("生花発送拠点 @川崎", "約500㎡の拠点に60㎡の花用冷蔵庫を完備。品質を保ったまま製作・発送できます。"),
         ("法人・大量発注に対応", "母の日などの繁忙期を含め、法人様向けの大量発注にも対応しています。")]
cw, ch, gap, top = (862 - 32) / 3, 100, 16, 372
for i, (t, sub) in enumerate(cards):
    x = 56 + i * (cw + gap)
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(top), Pt(cw), Pt(ch))
    r.adjustments[0] = 6 / ch
    r.fill.solid(); r.fill.fore_color.rgb = PINK
    r.line.fill.background()
    r.shadow.inherit = False
    st = r._element.find(qn("p:style"))
    if st is not None:
        r._element.remove(st)
    pic(f"_card{i}.png", x + 12, top + 11, 78, 78)
    text(x + 102, top + 14, cw - 114, 20, [(t, 12.8, DARK, True, 0)], wrap=False)
    text(x + 102, top + 36, cw - 114, 56, [(sub, 10.6, SUB, False, 0)], line=1.35)

# ---- 脚注 ----
text(56, 486, 862, 16, [("※ このほか、観葉植物EC「AND PLANTS」、花卉産業向け卸マーケットプレイス「ハナイチ」を運営しています。",
                         10.8, BROWN, False, 0.3)], wrap=False)

prs.save(OUT)
print("saved", OUT)

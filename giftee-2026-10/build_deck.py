"""ギフティ様向け 事業エコノミクス資料 — Domuz デザインシステム v2.1「発表スライド」型で .pptx を生成する。

retreat-2026-10/build_deck.py と同じ型（4:3・1128×846px 基準、Google スライドへは Drive の pptx 変換で取り込む）。

データの出どころ:
- AP/AF/ハナイチの月次（売上・限界利益・ROAS）: BigQuery and-plants の事業別月次エコノミクス（注文日ベース・税抜）
  26年6月までは「直近4期 月次エコノミクス」の広告費込み限界利益、26年7〜9月は 限界利益 − 売上÷ROAS で算出
- 現在のエコノミクス（AP/AF）: #リーダー 26年9月 貢献利益（速報値、AP Marketing Master ベース）
- 川崎拠点: スプレッドシート「花拠点_坪単位売上効率_月次推移」
"""
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu

# ---- デザイントークン（design-system/domuz_design_system.md §2） ----
MOSS = RGBColor(0x00, 0x43, 0x47)        # --brand-primary / --text-primary
LIGHT = RGBColor(0x00, 0xD7, 0x9C)       # --brand-secondary
FLOWER = RGBColor(0xEF, 0x66, 0x7D)      # --brand-flower
SECONDARY = RGBColor(0x62, 0x62, 0x62)   # --text-secondary
STRONG = RGBColor(0xE1, 0x64, 0x64)      # --text-strong
MUTED = RGBColor(0x99, 0x99, 0x99)       # --text-muted
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG_GRAY = RGBColor(0xCB, 0xD0, 0xD3)     # --bg-gray
BORDER_LIGHT = RGBColor(0xC9, 0xCF, 0xCF)
WATERMARK = RGBColor(0xEC, 0xEE, 0xEE)   # --watermark
PALE = RGBColor(0xE4, 0xE8, 0xE9)        # 売上の棒（背景寄りのグレー）

FONT_EN = "Montserrat"  # ギフティ向け株主報告資料と同じ
FONT_JA = "Noto Sans CJK JP"  # 和文は Montserrat にないため Noto で補う

W_PX, H_PX = 1128, 846
EMU_PER_PX = 9144000 / W_PX  # 10in 幅


def px(v):
    return Emu(int(round(v * EMU_PER_PX)))


def pt(px_size):
    """発表スライドの px サイズ → pt（38px ≒ 24pt）。"""
    return px_size * 24 / 38


prs = Presentation()
prs.slide_width = px(W_PX)
prs.slide_height = px(H_PX)
BLANK = prs.slide_layouts[6]
page_no = 0


def _set_font(run, size_px, color, bold, en=None):
    """英数字は Montserrat、和文は Noto Sans CJK JP（1つのランの中でも文字ごとに使い分けられる）。"""
    f = run.font
    f.size = Emu(int(pt(size_px) * 12700))
    f.bold = bold
    f.color.rgb = color
    f.name = FONT_EN
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("a:ea", FONT_JA), ("a:cs", FONT_EN)):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", face)


def text(slide, x, y, w, h, lines, size=25, color=MOSS, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.5, en=None, letter=None, para_gap=0):
    """lines: 文字列 or 文字列のリスト。[[語]] はライトグリーン、!!語!! は赤、**語** は太字。"""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if para_gap and i > 0:
            p.space_before = Emu(int(pt(para_gap) * 12700))
        for part in re.split(r"(\[\[.*?\]\]|\*\*.*?\*\*|!!.*?!!)", line):
            if not part:
                continue
            c, b = color, bold
            if part.startswith("[["):
                part, c, b = part[2:-2], LIGHT, True
            elif part.startswith("!!"):
                part, c, b = part[2:-2], STRONG, True
            elif part.startswith("**"):
                part, b = part[2:-2], True
            r = p.add_run()
            r.text = part
            _set_font(r, size, c, b, en)
            if letter:
                r.font._rPr.set("spc", str(int(letter * 100)))
    return tb


def _drop_style(shape):
    """テーマ由来の影（effectRef）を出さないよう、図形の p:style を外す。"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def rect(slide, x, y, w, h, fill, line=None, dash=False, shape=MSO_SHAPE.RECTANGLE, weight=1.5):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Emu(int(weight * 12700))
        if dash:
            s.line.dash_style = 4
    s.shadow.inherit = False
    _drop_style(s)
    return s


def hline(slide, x, y, w, color=BORDER_LIGHT, weight=1.0, dash=False):
    ln = slide.shapes.add_connector(1, px(x), px(y), px(x + w), px(y))
    ln.line.color.rgb = color
    ln.line.width = Emu(int(weight * 12700))
    if dash:
        ln.line.dash_style = 4
    _drop_style(ln)
    return ln


def vline(slide, x, y, h, color=BORDER_LIGHT, weight=1.0, dash=False):
    ln = slide.shapes.add_connector(1, px(x), px(y), px(x), px(y + h))
    ln.line.color.rgb = color
    ln.line.width = Emu(int(weight * 12700))
    if dash:
        ln.line.dash_style = 4
    _drop_style(ln)
    return ln


def polyline(slide, pts, color, weight=2.5):
    """pts: [(x_px, y_px), ...] を折れ線で描く（Google スライドでも図形として残る）。"""
    fb = slide.shapes.build_freeform(px(pts[0][0]), px(pts[0][1]), scale=1.0)
    fb.add_line_segments([(px(x), px(y)) for x, y in pts[1:]], close=False)
    s = fb.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Emu(int(weight * 12700))
    s.shadow.inherit = False
    _drop_style(s)
    return s


def badge(slide, x, y, n, color=MOSS, d=22):
    """トピック番号の丸。"""
    rect(slide, x - d / 2, y - d / 2, d, d, color, shape=MSO_SHAPE.OVAL)
    text(slide, x - d / 2, y - d / 2, d, d, str(n), size=13, color=WHITE, bold=True, en=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)


def content_slide(title=None, sub=None):
    """白背景＋左端ライトグリーン縦帯8px＋左上タイトル＋右下ページ番号。"""
    global page_no
    page_no += 1
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 8, H_PX, LIGHT)
    if title:
        text(s, 56, 44, 1000, 56, title, size=38, bold=True)
    if sub:
        text(s, 56, 104, 1000, 36, sub, size=20, color=SECONDARY)
    text(s, W_PX - 120, H_PX - 44, 80, 24, str(page_no), size=14, color=MUTED,
         align=PP_ALIGN.RIGHT, en=True)
    return s


def footnote(s, body):
    text(s, 56, H_PX - 84, 960, 40, body, size=13, color=SECONDARY, spacing=1.3)


def divider(num, en_label, title, sub=""):
    """P-3 章扉: Moss 全面＋アウトライン章番号＋英字 ALL CAPS＋和文タイトル。"""
    global page_no
    page_no += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = MOSS
    tb = text(s, 80, 240, 400, 150, num, size=120, color=WHITE, en=True, spacing=1.0)
    rpr = tb.text_frame.paragraphs[0].runs[0].font._rPr
    for el in rpr.findall(qn("a:solidFill")):
        rpr.remove(el)
    ln = rpr.makeelement(qn("a:ln"), {"w": "19050"})
    sf = ln.makeelement(qn("a:solidFill"), {})
    sf.append(sf.makeelement(qn("a:srgbClr"), {"val": "FFFFFF"}))
    ln.append(sf)
    rpr.insert(0, ln)
    rpr.insert(1, rpr.makeelement(qn("a:noFill"), {}))
    text(s, 80, 416, 900, 32, en_label, size=22, color=WHITE, en=True, letter=6)
    text(s, 80, 460, 960, 64, title, size=46, color=WHITE, bold=True)
    if sub:
        text(s, 80, 544, 960, 36, sub, size=22, color=WHITE)
    return s


# =====================================================================
# データ（2022年10月〜2026年9月の48か月。単位：百万円・税抜）
# =====================================================================
YM = [(2022 + (9 + i) // 12, (9 + i) % 12 + 1) for i in range(48)]  # (年, 月)
FY_NAMES = ["23年9月期", "24年9月期", "25年9月期", "26年9月期"]

AP_S = [11.2, 10.6, 10.5, 12.0, 13.1, 16.4, 19.8, 19.1, 19.1, 24.8, 19.9, 17.7,
        22.2, 20.3, 22.5, 29.3, 24.7, 36.6, 37.6, 40.0, 40.1, 35.9, 37.2, 46.5,
        39.3, 33.2, 37.5, 59.0, 39.3, 47.2, 61.9, 62.3, 46.8, 47.5, 38.8, 42.0,
        41.3, 39.4, 37.5, 63.1, 34.5, 43.6, 58.6, 65.5, 47.9, 49.0, 40.1, 42.6]
AP_MP = [5.4, 5.0, 5.0, 5.6, 6.1, 7.8, 9.4, 9.2, 9.3, 11.4, 8.9, 8.1,
         9.9, 8.8, 10.4, 13.1, 11.5, 16.3, 17.2, 18.4, 18.4, 16.3, 17.1, 21.0,
         17.9, 15.4, 18.1, 25.9, 18.8, 21.5, 29.0, 29.0, 21.5, 21.2, 17.9, 19.7,
         18.6, 18.5, 17.8, 27.8, 16.8, 20.7, 27.8, 31.1, 22.5, 22.3, 18.3, 19.6]
AP_NET = [3.3, 2.7, 2.9, 3.7, 3.6, 5.1, 5.6, 4.7, 6.5, 5.8, 4.6, 4.0,
          4.6, 4.0, 5.9, 8.5, 6.5, 9.4, 8.8, 7.3, 10.1, 8.0, 8.6, 11.1,
          10.0, 9.1, 8.3, 18.3, 11.6, 15.7, 16.5, 12.7, 11.9, 11.7, 11.6, 13.2,
          11.6, 12.0, 12.0, 21.5, 13.6, 17.5, 17.1, 17.4, 14.2]
AP_ROAS = [545, 478, 506, 634, 526, 605, 514, 422, 669, 440, 467, 427,
           422, 418, 499, 626, 493, 530, 450, 360, 482, 432, 435, 468,
           495, 529, 381, 785, 544, 810, 499, 383, 491, 497, 616, 644,
           588, 605, 651, 996, 1061, 1342, 544, 478, 579, 564, 543, 613]

AF_S = [1.1, 1.1, 1.9, 1.7, 1.7, 2.8, 6.4, 8.7, 1.8, 2.3, 2.2, 3.7,
        3.3, 3.7, 3.9, 4.5, 5.6, 10.1, 16.3, 19.6, 5.6, 6.9, 9.3, 10.8,
        11.0, 10.3, 11.8, 9.9, 12.6, 14.0, 21.7, 23.8, 10.8, 11.2, 10.5, 13.9,
        14.6, 14.6, 11.9, 12.1, 10.4, 14.6, 27.8, 30.3, 8.4, 7.5, 9.0, 11.3]
AF_MP = [0.6, 0.5, 1.0, 0.8, 0.8, 1.4, 2.6, 3.8, 0.9, 1.1, 1.1, 1.7,
         1.6, 1.8, 1.9, 2.0, 2.5, 4.4, 8.0, 10.4, 2.5, 2.9, 4.0, 4.7,
         4.8, 4.6, 5.3, 4.3, 5.6, 6.3, 11.0, 12.6, 4.9, 5.1, 4.7, 6.3,
         6.4, 6.6, 5.3, 5.2, 4.6, 6.7, 15.5, 17.7, 4.1, 3.5, 4.3, 5.5]
AF_NET = [0.5, 0.3, 0.7, 0.4, 0.3, 0.6, 1.4, 2.5, 0.0, 0.0, 0.4, 0.4,
          0.6, 0.6, 0.9, 0.8, 1.1, 1.5, 3.1, 5.4, 0.0, 0.0, -0.2, 0.1,
          -1.1, -0.6, 0.2, 0.2, -0.8, 1.1, 4.9, 7.8, -0.3, -0.4, 0.7, 2.3,
          1.9, 2.0, 1.1, 1.5, 2.3, 3.5, 11.9, 13.7, 2.0]
AF_ROAS = [1868, 627, 538, 415, 315, 373, 521, 665, 218, 201, 347, 288,
           327, 302, 424, 378, 385, 335, 335, 398, 228, 235, 219, 237,
           184, 196, 229, 240, 196, 270, 360, 487, 212, 206, 266, 346,
           331, 316, 287, 320, 438, 464, 763, 755, 408, 445, 480, 376]

# ハナイチは 2024年2月から（それ以前は 0）
_pad = [0.0] * 16
HA_S = _pad + [0.0, 0.63, 1.49, 0.92, 1.66, 3.44, 2.05, 1.86, 1.44, 0.45, 1.03, 0.6, 1.76, 3.37, 4.64, 2.49,
               3.64, 3.94, 3.57, 3.45, 2.3, 2.43, 2.72, 1.88, 2.34, 5.64, 5.79, 4.42, 3.64, 4.19, 5.94, 5.53]
HA_GP = _pad + [0.0, 0.26, 0.59, 0.37, 0.67, 1.36, 0.81, 0.71, 0.56, 0.14, 0.27, 0.17, 0.47, 0.97, 1.33, 0.74,
                1.09, 1.34, 1.12, 1.0, 0.57, 0.63, 0.76, 0.6, 0.76, 2.06, 2.17, 1.64, 1.28, 1.39, 1.71, 1.52]
HA_ROAS = [None] * 25 + [5989, 1244, 577, 491, 346, 295, 255, 215, 453, 570, 688, 381, 387, 448, 364, 545,
                         1249, 1766, 1690, 1264, 727, 448, 398]


def ad_from_roas(sales, roas):
    return [s * 100 / r if r else 0.0 for s, r in zip(sales, roas)]


AP_AD = ad_from_roas(AP_S, AP_ROAS)
AF_AD = ad_from_roas(AF_S, AF_ROAS)
HA_AD = ad_from_roas(HA_S, HA_ROAS)
# 26年7〜9月の広告費込み限界利益 = 限界利益 − 広告費
AP_NET = AP_NET + [m - a for m, a in zip(AP_MP[45:], AP_AD[45:])]
AF_NET = AF_NET + [m - a for m, a in zip(AF_MP[45:], AF_AD[45:])]
HA_NET = [g - a for g, a in zip(HA_GP, HA_AD)]
assert all(len(v) == 48 for v in (AP_S, AP_MP, AP_NET, AP_AD, AF_S, AF_MP, AF_NET, AF_AD, HA_S, HA_GP, HA_NET, HA_AD))


def fy_sum(series):
    return [sum(series[i * 12:(i + 1) * 12]) for i in range(4)]


# =====================================================================
# 1. 表紙
# =====================================================================
page_no += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 8, H_PX, LIGHT)
text(s, 96, 250, 960, 40, "株式会社ギフティ様　ご参考資料", size=26, color=SECONDARY)
text(s, 96, 330, 980, 100, "事業エコノミクスの整理とまとめ", size=60, bold=True, spacing=1.3)
text(s, 96, 450, 960, 40, "AP/AF・ハナイチの推移、今後のエコノミクス、川崎拠点の坪効率", size=24,
     color=SECONDARY)
text(s, 96, 700, 900, 32, "2026年10月　株式会社DOMUZ　高木裕樹", size=22, color=SECONDARY)

# =====================================================================
# 2. 本資料の構成・要約・用語
# =====================================================================
s = content_slide("本資料の構成")
items = [
    ("1", "各事業のエコノミクスの整理と経年推移", "アンドプランツ/アンドフラワー、ハナイチ（23年9月期〜26年9月期）"),
    ("2", "今後の事業におけるエコノミクスの整理", "閑散期・繁忙期の原価・広告ROAS・直接固定費と、Value UPプラン"),
    ("3", "倉庫としての収益性の試算", "川崎拠点の坪効率（単純3PL／高付加価値3PL／AP・AF）"),
]
y = 200
for n, label, sub in items:
    text(s, 96, y, 80, 104, n, size=64, color=LIGHT, bold=True, en=True, anchor=MSO_ANCHOR.MIDDLE,
         spacing=1.0)
    text(s, 200, y + 8, 860, 50, label, size=34, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    text(s, 200, y + 64, 860, 32, sub, size=20, color=SECONDARY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    if n != "3":
        hline(s, 96, y + 128, 936)
    y += 160
footnote(s, "※ 用語と集計の前提は4ページにまとめています")

summary_slide = content_slide("要約")
terms = content_slide("用語と集計の前提")
TERMS = [
    ("期", "9月決算。「26年9月期」＝2025年10月〜2026年9月"),
    ("売上", "税抜・注文日ベース。アンドプランツ/アンドフラワーは自社EC＋モール（楽天・Amazon・Yahoo!）"),
    ("限界利益", "売上 − 原価・送料・資材・決済手数料（ハナイチは粗利＝売上 − 仕入原価）"),
    ("広告費込み限界利益", "限界利益 − 広告費。AP/AFに共通する広告（母の日・指名検索など）は売上比で配賦"),
    ("ROAS", "売上 ÷ 広告費"),
    ("直接固定費", "事業に直接かかる人件費（アルバイト・スポット・外部加工）と、水道光熱・車両・修繕・消耗品・通信・賃借料"),
    ("貢献利益", "広告費込み限界利益 − 直接固定費。本社の人件費・共通費は含まない"),
    ("出典", "受注データ（BigQuery）、社内の貢献利益集計（月次速報値を含む）、花拠点の坪効率シート。試算表とは集計基準が異なる"),
]
y = 160
for k_, v_ in TERMS:
    text(terms, 56, y, 250, 60, k_, size=19, bold=True)
    text(terms, 300, y, 772, 60, v_, size=17, color=SECONDARY, spacing=1.4)
    y += 72
    hline(terms, 56, y - 14, 1016)

# =====================================================================
# 章 01
# =====================================================================
divider("01", "ECONOMICS TREND", "各事業のエコノミクスの推移", "23年9月期（2022年10月）〜 26年9月期（2026年9月）")


def trend_slide(title, sales, mp, ad, net, color, ymax, note, msg, fy0=0, mp_label="限界利益",
                net_label="広告費込み限界利益", fy_names=None):
    """月次売上の棒（事業の色）＋期別の収益性の表。棒の上は各期の最高月、期名の下は期合計。"""
    s = content_slide(title)
    if msg:
        text(s, 56, 104, 1016, 36, msg, size=21, bold=True)
    fy_names = fy_names or FY_NAMES
    x0, x1 = 214, 914
    text(s, x0, 154, 900, 24, f"{FY_NAMES[fy0]}（{YM[fy0 * 12][0]}年10月）〜26年9月期（2026年9月）　単位：百万円　"
         "棒の上は各期の最高月、期名の下は期合計", size=14, color=SECONDARY)
    top, base = 206, 434
    n = 48 - fy0 * 12
    slot = (x1 - x0) / n
    bw = slot * 0.66
    vmax = max(sales)
    sy = lambda v: base - (base - top) * v / vmax
    for k in range(fy0, 4):
        idx = range(k * 12, k * 12 + 12)
        peak = max(idx, key=lambda i: sales[i])
        for i in idx:
            if sales[i] <= 0:
                continue
            bx = x0 + (i - fy0 * 12) * slot + (slot - bw) / 2
            rect(s, bx, sy(sales[i]), bw, base - sy(sales[i]), color)
            if i == peak:
                text(s, bx - 20, sy(sales[i]) - 26, bw + 40, 22, f"{sales[i]:.0f}" if vmax >= 20 else f"{sales[i]:.1f}", size=15, bold=True,
                     align=PP_ALIGN.CENTER, en=True)
        gx = x0 + (k - fy0) * 12 * slot
        if k > fy0:
            vline(s, gx, 190, base - 190 + 70, BORDER_LIGHT, 1.0, dash=True)
        text(s, gx, base + 12, 12 * slot, 26, fy_names[k], size=17, bold=True, align=PP_ALIGN.CENTER)
        text(s, gx, base + 40, 12 * slot, 24, f"{sum(sales[i] for i in idx):.0f}", size=16, color=SECONDARY,
             align=PP_ALIGN.CENTER, en=True)
    hline(s, x0 - 4, base, x1 - x0 + 8, BORDER_LIGHT, 1.5)
    # 期別の表
    hdr = ["期", "売上", mp_label, f"{mp_label}率", "広告費", "全体ROAS", net_label, "同率"]
    cxs = [56, 190, 290, 410, 520, 620, 740, 950]
    cws = [130, 100, 120, 110, 100, 120, 210, 122]
    ty = 512
    rect(s, 56, ty, 1016, 40, MOSS)
    for j, h in enumerate(hdr):
        text(s, cxs[j] + (12 if j == 0 else 0), ty, cws[j] - (0 if j == 0 else 12), 40, h, size=16, color=WHITE,
             bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
    y = ty + 40
    fs, fm, fa, fn = fy_sum(sales), fy_sum(mp), fy_sum(ad), fy_sum(net)
    rh = 48 if fy0 == 0 else 56
    for k in range(fy0, 4):
        if k == 3:
            rect(s, 56, y, 1016, rh, WATERMARK)
        roas = f"{fs[k] / fa[k] * 100:.0f}%" if fa[k] > 0 else "—"
        vals = [fy_names[k], f"{fs[k]:.1f}", f"{fm[k]:.1f}", f"{fm[k] / fs[k] * 100:.1f}%", f"{fa[k]:.1f}" if fa[k] > 0 else "—", roas,
                f"{fn[k]:.1f}", f"{fn[k] / fs[k] * 100:.1f}%"]
        for j, v in enumerate(vals):
            text(s, cxs[j] + (12 if j == 0 else 0), y, cws[j] - (0 if j == 0 else 12), rh, v, size=19, bold=True,
                 color=STRONG if v.startswith("-") else MOSS, anchor=MSO_ANCHOR.MIDDLE,
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
        y += rh
        hline(s, 56, y, 1016, BORDER_LIGHT)
    footnote(s, note)
    return s


ap_fy_s, ap_fy_net, ap_fy_ad = fy_sum(AP_S), fy_sum(AP_NET), fy_sum(AP_AD)
af_fy_s, af_fy_net, af_fy_ad = fy_sum(AF_S), fy_sum(AF_NET), fy_sum(AF_AD)
ha_fy_s, ha_fy_net, ha_fy_ad, ha_fy_gp = fy_sum(HA_S), fy_sum(HA_NET), fy_sum(HA_AD), fy_sum(HA_GP)

AP_GREEN = LIGHT
trend_slide(
    "アンドプランツ：月次売上と期別の収益性", AP_S, AP_MP, AP_AD, AP_NET, AP_GREEN, 70,
    "※ 税抜・注文日ベース。自社EC＋モール（楽天・Amazon・Yahoo!）。BigQuery（and-plants）の事業別月次。"
    "モール広告費は25年6月以降のみ計上（それ以前の広告費込み限界利益は実態より高め）",
    f"広告費を絞って広告費込み限界利益は{ap_fy_net[3]:.0f}百万円（前期比 +{(ap_fy_net[3] / ap_fy_net[2] - 1) * 100:.0f}%）に",
)
trend_slide(
    "アンドフラワー：月次売上と期別の収益性", AF_S, AF_MP, AF_AD, AF_NET, FLOWER, 35,
    "※ 税抜・注文日ベース。自社EC（モールの花は微小のため含む）。BigQuery（and-plants）の事業別月次",
    "",
)
trend_slide(
    "ハナイチ：月次売上と期別の収益性", HA_S, HA_GP, HA_AD, HA_NET, MOSS, 7,
    "※ 税抜・注文日ベース。ハナイチ注文DB（BigQuery）の注文合計で、卸・法人など注文DB外の売上は含まない。"
    "24年9月期は2024年3月〜9月の7か月。広告は24年11月に開始",
    f"売上は{ha_fy_s[3]:.0f}百万円（前期比 +{(ha_fy_s[3] / ha_fy_s[2] - 1) * 100:.0f}%）。注文数が伸び、広告費込み粗利も改善",
    fy0=1, mp_label="粗利", net_label="広告費込み粗利",
)

# =====================================================================
# 主なトピック
# =====================================================================
s = content_slide("各事業の主なトピック", "23年9月期〜26年9月期")
TOPICS = [
    ("アンドプランツ", AP_GREEN, [
        ("25年1月", "月商59百万円（当時の過去最高）"),
        ("25年6月〜", "モール（Amazon・楽天）の広告費も計上"),
        ("26年1月", "月商63百万円"),
        ("26年2〜3月", "広告を絞り、ROAS1,000%超で利益を維持"),
        ("26年5月", "月商65.5百万円（過去最高）")]),
    ("アンドフラワー", FLOWER, [
        ("毎年5月", "母の日。26年5月は月商30.3百万円"),
        ("24年10月〜25年3月", "ROAS200%前後で広告費込み限界利益が赤字圏"),
        ("25年3月", "川崎拠点を開設"),
        ("26年4月", "新城拠点を閉じ、川崎に集約"),
        ("26年5月", "広告費込み限界利益13.7百万円（過去最高）")]),
    ("ハナイチ", MOSS, [
        ("24年3月", "サービス開始"),
        ("24年11月", "広告を開始"),
        ("25年3月", "生花の注文が増える（月133件）"),
        ("26年3月", "月商5.6百万円（前年同月比 +67%）"),
        ("26年8月", "注文430件（前年同月の3.5倍）")]),
]
for k, (name, col, items_) in enumerate(TOPICS):
    x = 56 + k * 344
    rect(s, x, 168, 320, 6, col)
    text(s, x, 186, 320, 36, name, size=26, bold=True, color=FLOWER if col == FLOWER else MOSS)
    yy = 244
    for when, body in items_:
        text(s, x, yy, 320, 26, when, size=17, bold=True, color=SECONDARY)
        text(s, x, yy + 28, 320, 60, body, size=17, spacing=1.3)
        yy += 98
text(s, 56, 748, 1016, 40, "APは[[広告効率の改善で利益が残る形]]に。AFは母の日、ハナイチは注文数の伸びが柱", size=22, bold=True)

# =====================================================================
# 章 02
# =====================================================================
divider("02", "UNIT ECONOMICS", "今後の事業におけるエコノミクス", "原価・広告ROAS・直接固定費")

# ---- 現在のエコノミクス（閑散期＝26年9月、繁忙期＝26年5月）
# AP/AF: 社内の貢献利益集計（9月は #リーダー の速報値、5月は「貢献利益（共有用）」）。
# 広告費 = 売上 × 限界利益率（BigQuery の同月） − 広告費込み限界利益（共通広告は売上比で配賦されたもの）
def econ(S, NET, LABOR, CONTRIB, i, mp_series, s_series):
    mpr = mp_series[i] / s_series[i]
    e = dict(S=S, NET=NET, LABOR=LABOR, CONTRIB=CONTRIB)
    e["COGS"] = S * (1 - mpr)
    e["AD"] = S * mpr - NET
    e["ROAS"] = S / e["AD"] * 100
    e["OTHER"] = NET - LABOR - CONTRIB
    return e


def ha_econ(i):
    e = dict(S=HA_S[i], GP=HA_GP[i], ROAS=HA_ROAS[i])
    e["AD"] = e["S"] * 100 / e["ROAS"]
    e["COGS"] = e["S"] - e["GP"]
    e["NET"] = e["GP"] - e["AD"]
    return e


PERIODS = {
    "sep": dict(label="閑散期（26年9月）",
                ap=econ(44.44, 12.77, 2.41, 5.27, 47, AP_MP, AP_S),
                af=econ(11.35, 2.57, 3.28, -4.07, 47, AF_MP, AF_S),
                ha=ha_econ(47)),
    "may": dict(label="繁忙期（26年5月・母の日）",
                ap=econ(65.90, 21.70, 3.59, 12.73, 43, AP_MP, AP_S),
                af=econ(30.41, 9.42, 6.09, -0.36, 43, AF_MP, AF_S),
                ha=ha_econ(43)),
}
ap, af = PERIODS["sep"]["ap"], PERIODS["sep"]["af"]
HA_NOW = PERIODS["sep"]["ha"]


def waterfall(s, x, name, steps, end_label, end_val, sales, color, lo=-40):
    """steps: [(ラベル, 金額)] を売上から順に引く。縦軸は 100 〜 lo（%）で3事業共通。"""
    text(s, x, 156, 330, 34, name, size=24, bold=True, color=color)
    text(s, x, 192, 330, 26, f"売上 {sales:.1f}百万円", size=17, color=SECONDARY)
    top, bot = 262, 612
    yv = lambda v: top + (bot - top) * (100 - v) / (100 - lo)
    bw, gap = 46, 12
    bx = x
    rect(s, bx, yv(100), bw, yv(0) - yv(100), PALE)
    text(s, bx - 10, yv(100) - 24, bw + 20, 22, "100", size=14, bold=True, align=PP_ALIGN.CENTER, en=True)
    labels = [(bx, "売上", SECONDARY, False)]
    cur = 100.0
    for lab, v in steps:
        bx += bw + gap
        p = v / sales * 100
        y0, y1 = yv(cur), yv(cur - p)
        rect(s, bx, y0, bw, y1 - y0, BG_GRAY)
        text(s, bx - 10, y1 + 3, bw + 20, 20, f"{p:.0f}", size=13, color=SECONDARY, align=PP_ALIGN.CENTER, en=True)
        labels.append((bx, lab, SECONDARY, False))
        cur -= p
    bx += bw + gap
    endp = end_val / sales * 100
    col = LIGHT if endp >= 0 else STRONG
    y0, y1 = sorted((yv(0), yv(endp)))
    rect(s, bx, y0, bw, max(y1 - y0, 2), col)
    ly = y0 - 24 if endp >= 0 else y1 + 3
    text(s, bx - 6, ly, bw + 12, 22, f"{endp:.0f}", size=15, bold=True, color=col, align=PP_ALIGN.CENTER, en=True)
    labels.append((bx, end_label, MOSS, True))
    hline(s, x - 6, yv(0), bx + bw + 12 - x, MUTED)
    for lx, lab, c, b in labels:
        text(s, lx - 10, bot + 8, bw + 20, 44, lab, size=12, color=c, bold=b, align=PP_ALIGN.CENTER, spacing=1.2)
    return s


def econ_slide(key, n, takeaway, note):
    P = PERIODS[key]
    s = content_slide(f"エコノミクス{n}：{P['label']}", "売上を100としたときの内訳　※税抜")
    a_, f_, h_ = P["ap"], P["af"], P["ha"]
    waterfall(s, 56, "アンドプランツ",
              [("原価等", a_["COGS"]), ("広告費", a_["AD"]), ("人件費", a_["LABOR"]), ("その他", a_["OTHER"])],
              "貢献\n利益", a_["CONTRIB"], a_["S"], MOSS)
    waterfall(s, 420, "アンドフラワー",
              [("原価等", f_["COGS"]), ("広告費", f_["AD"]), ("人件費", f_["LABOR"]), ("その他", f_["OTHER"])],
              "貢献\n利益", f_["CONTRIB"], f_["S"], FLOWER)
    waterfall(s, 784, "ハナイチ", [("原価", h_["COGS"]), ("広告費", h_["AD"])],
              "広告費\n込み粗利", h_["NET"], h_["S"], MOSS)
    text(s, 784 + 3 * 58 + 4, 300, 150, 120, ["人件費などの", "直接固定費は", "確認中"], size=14, color=MUTED, spacing=1.3)
    vline(s, 400, 160, 500)
    vline(s, 764, 160, 500)
    text(s, 56, 678, 1016, 60, takeaway, size=20, bold=True, spacing=1.4)
    footnote(s, note)


NOTE_ECON = ("広告費は受注データの限界利益率から算出（共通広告は売上比で配賦）。その他＝光熱・消耗品・賃借料等（貢献利益からの逆算）。"
             "ハナイチは受注データの粗利")
econ_slide("sep", "①", ["アンドプランツは平常月でも貢献利益が残る。アンドフラワーは原価・広告は回っているが、",
                        "[[固定費を賄う売上規模]]に届いていない"],
           "※ 社内の9月貢献利益（速報値。APはSBSの棚卸費用を含む）。" + NOTE_ECON)
econ_slide("may", "②", ["繁忙期はアンドプランツの貢献利益が月12.7百万円。アンドフラワーは売上が3倍になるが、",
                        "人件費（スポット含む）も増え、[[母の日の月でほぼ損益分岐]]"],
           "※ 社内の貢献利益集計（26年5月）。APの人件費は外部加工費1.1百万円を含む。" + NOTE_ECON)

# ---- 主要指標（表）：閑散期と繁忙期を並べる
s = content_slide("エコノミクスの主要指標", "閑散期（26年9月）と繁忙期（26年5月・母の日）の比較　※税抜")
cx2 = [56, 336, 458, 580, 702, 824, 946]
y = 160
for j, (name, col) in enumerate([("アンドプランツ", MOSS), ("アンドフラワー", FLOWER), ("ハナイチ", MOSS)]):
    text(s, cx2[1 + 2 * j], y, 240, 28, name, size=18, bold=True, color=col, align=PP_ALIGN.CENTER)
    hline(s, cx2[1 + 2 * j] + 12, y + 32, 220, col, 2.0)
    for q, lab in enumerate(["閑散期", "繁忙期"]):
        text(s, cx2[1 + 2 * j + q], y + 40, 114, 24, lab, size=15, color=SECONDARY, align=PP_ALIGN.RIGHT)
hline(s, 56, y + 74, 1016, MOSS, 1.5)
y += 88
SEQ = [PERIODS["sep"], PERIODS["may"]]


def r2(label, fn, sub=None, bold=False):
    global y
    text(s, cx2[0], y, 280, 30, label, size=18, bold=True)
    if sub:
        text(s, cx2[0], y + 28, 280, 22, sub, size=12, color=SECONDARY)
    vals = []
    for b_ in ("ap", "af", "ha"):
        for P in SEQ:
            vals.append(fn(b_, P[b_]))
    for j, v in enumerate(vals):
        text(s, cx2[1 + j], y, 114, 30, v, size=18, bold=bold, align=PP_ALIGN.RIGHT,
             color=STRONG if v.startswith("▲") else (MUTED if v.startswith("確認") else MOSS))
    y += 62 if sub else 48
    hline(s, 56, y - 12, 1016)


def yen(v):
    return f"▲{-v:.1f}" if v < 0 else f"{v:.1f}"


r2("売上（月・百万円）", lambda b, e: f"{e['S']:.1f}")
r2("原価率", lambda b, e: f"{e['COGS'] / e['S'] * 100:.0f}%", "AP/AFは原価・送料・資材・決済手数料を含む")
r2("広告ROAS", lambda b, e: f"{e['ROAS']:.0f}%", "売上÷広告費")
r2("広告費込み限界利益率", lambda b, e: f"{e['NET'] / e['S'] * 100:.1f}%", "ハナイチは広告費込み粗利率", bold=True)
r2("直接固定費（月・百万円）", lambda b, e: "確認中" if b == "ha" else f"{e['LABOR'] + e['OTHER']:.1f}",
   "人件費＋光熱・消耗品・賃借料等")
r2("貢献利益（月・百万円）", lambda b, e: "確認中" if b == "ha" else yen(e["CONTRIB"]), bold=True)
af_be = (af["NET"] - af["CONTRIB"]) / (af["NET"] / af["S"])
text(s, 56, y + 6, 1016, 70,
     [f"アンドフラワーの損益分岐は、閑散期の利益率・固定費で月商 約[[{af_be:.0f}百万円]]。",
      "繁忙期は利益率が上がる一方、人件費も増えるため、月商30百万円でほぼトントン"], size=19, bold=True, spacing=1.45)
footnote(s, "※ 損益分岐＝直接固定費÷広告費込み限界利益率。固定費が売上に比例して増えない前提の概算。ハナイチの直接固定費は確認中")

# ---- Value UP プラン
s = content_slide("今後のValue UPプラン", "数量が増えたときに打てる手 ── 卸と輸入で原価を下げ、売価を下げる")
cols3 = [("アンドプランツ", MOSS,
          ["観葉植物・鉢を生産者／海外から直接仕入れ（輸入）し、仕入単価を下げる",
           "卸（法人・小売向け）を並行し、仕入ロットを大きくする",
           "下がった原価の一部を売価に回し、CVRとROASを上げる"]),
         ("アンドフラワー", FLOWER,
          ["輸入花・産地直送を増やし、市場経由の仕入を減らす",
           "ハナイチの仕入網と仕入量をまとめて使い、単価を下げる",
           "平常月の売価を下げて数量を取り、固定費を薄める"]),
         ("ハナイチ", MOSS,
          ["取扱量の増加を背景に、産地・輸入元と直接取引する",
           "鉢・資材は卸ルートで調達し、粗利率を上げる",
           "小売店・教室向けに価格を下げ、継続の仕入先になる"])]
for k, (name, col, pts) in enumerate(cols3):
    x = 56 + k * 344
    rect(s, x, 168, 320, 6, col)
    text(s, x, 186, 320, 36, name, size=26, bold=True, color=col)
    yy = 240
    for p_ in pts:
        rect(s, x, yy + 10, 8, 8, LIGHT)
        text(s, x + 20, yy, 300, 110, p_, size=18, spacing=1.4)
        yy += 112
text(s, 56, 600, 1016, 80, ["数量 → 仕入条件の改善 → 原価↓ → 売価↓ → 数量↑ の[[循環をつくる]]",
                             "ギフティ様の販路での数量は、この循環の起点になる"], size=22, bold=True, spacing=1.5)
footnote(s, "※ 方向性の整理。具体的な仕入先・条件は今後詰める")

# ---- Value UP の効果（感応度）
s = content_slide("原価を下げたときの効果（試算）", "26年9月期の売上規模のまま、原価率だけが下がった場合　単位：百万円/年")
fy26 = [("アンドプランツ", ap_fy_s[3], MOSS), ("アンドフラワー", af_fy_s[3], FLOWER), ("ハナイチ", ha_fy_s[3], MOSS)]
cx3 = [56, 400, 600, 800]
y = 176
for j, c in enumerate(["", "26年9月期 売上", "原価率 −3pt", "原価率 −5pt"]):
    text(s, cx3[j], y, 200, 30, c, size=18, color=SECONDARY, bold=True,
         align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
hline(s, 56, y + 40, 1016, MOSS, 1.5)
y += 56
tot3 = tot5 = 0
for name, sv, col in fy26:
    text(s, cx3[0], y, 300, 30, name, size=20, bold=True, color=col)
    text(s, cx3[1], y, 200, 30, f"{sv:.0f}", size=21, align=PP_ALIGN.RIGHT, en=True)
    text(s, cx3[2], y, 200, 30, f"+{sv * 0.03:.1f}", size=21, bold=True, color=LIGHT, align=PP_ALIGN.RIGHT, en=True)
    text(s, cx3[3], y, 200, 30, f"+{sv * 0.05:.1f}", size=21, bold=True, color=LIGHT, align=PP_ALIGN.RIGHT, en=True)
    tot3 += sv * 0.03
    tot5 += sv * 0.05
    y += 48
    hline(s, 56, y - 10, 1016)
text(s, cx3[0], y, 300, 30, "合計", size=20, bold=True)
text(s, cx3[2], y, 200, 30, f"+{tot3:.1f}", size=22, bold=True, color=LIGHT, align=PP_ALIGN.RIGHT, en=True)
text(s, cx3[3], y, 200, 30, f"+{tot5:.1f}", size=22, bold=True, color=LIGHT, align=PP_ALIGN.RIGHT, en=True)
y += 70
m_now = af["NET"] / af["S"]
fixed = af["NET"] - af["CONTRIB"]
text(s, 56, y, 1016, 30, "アンドフラワーの損益分岐（月商）", size=22, bold=True)
y += 44
for lab, m in [("現状", m_now), ("原価率 −5pt", m_now + 0.05), ("原価率 −10pt", m_now + 0.10)]:
    text(s, 56, y, 300, 30, lab, size=19, color=SECONDARY)
    w = 560 * (fixed / m) / 30
    rect(s, 300, y + 4, w, 22, MOSS if lab == "現状" else LIGHT)
    text(s, 300 + w + 12, y, 200, 30, f"{fixed / m:.0f}百万円", size=19, bold=True, en=False)
    y += 40
footnote(s, "※ 原価率の改善分がそのまま利益になる前提（売価を下げない場合）。売価に回す分は数量の増加で回収する想定。"
         "損益分岐は直接固定費を26年9月の水準で固定した概算")

# =====================================================================
# 章 03
# =====================================================================
divider("03", "WAREHOUSE", "倉庫としての収益性", "川崎拠点の坪効率を、用途別に見える化する")

# ---- 川崎拠点の現状
KAWA_TSUBO, KAWA_RENT = 128.84, 0.94  # 坪, 百万円/月
WH = [  # (ラベル, 面積, AF限界利益, ハナイチ限界利益)  単位：百万円。25年3月〜26年3月は新城40.49坪を含む
    ("25/6", 169.33, 4.851, 0.0), ("7", 169.33, 5.077, 0.0), ("8", 169.33, 4.668, 0.051), ("9", 169.33, 6.309, 0.052),
    ("10", 169.33, 6.357, 0.127), ("11", 169.33, 6.573, -0.019), ("12", 169.33, 5.258, 0.0), ("26/1", 169.33, 5.212, 0.140),
    ("2", 169.33, 4.622, 0.162), ("3", 169.33, 6.654, 0.245), ("4", 128.84, 15.393, 0.205), ("5", 128.84, 17.552, 0.315),
    ("6", 128.84, 4.043, 0.120), ("7", 128.84, 3.513, 0.165), ("8", 128.84, 4.289, 0.450),
]
s = content_slide("川崎拠点の現状", "花の拠点の坪あたり限界利益（月）　単位：千円/坪")
kp = [("面積", f"{KAWA_TSUBO:.2f}", "坪"), ("家賃", "94", "万円/月"), ("坪単価", "7,296", "円/坪")]
for k, (lab, v, u) in enumerate(kp):
    x = 56 + k * 200
    text(s, x, 150, 190, 26, lab, size=17, color=SECONDARY)
    tb = text(s, x, 176, 190, 50, v, size=40, bold=True, en=True, spacing=1.0)
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = u
    _set_font(r, 17, MOSS, True, en=False)
# 棒（AF + ハナイチの積み上げ）と家賃の坪単価
x0, top, base = 112, 290, 600
slot = (1060 - x0) / len(WH)
bw = slot * 0.6
ymax = 140
sy = lambda v: base - (base - top) * v / ymax
for v in range(0, ymax + 1, 20):
    hline(s, x0 - 4, sy(v), 1060 - x0 + 4, BORDER_LIGHT if v == 0 else WATERMARK)
    text(s, 56, sy(v) - 11, 46, 22, str(v), size=13, color=MUTED, align=PP_ALIGN.RIGHT, en=True)
for i, (lab, area, af_mp, ha_mp) in enumerate(WH):
    bx = x0 + i * slot + (slot - bw) / 2
    a = af_mp * 1000 / area
    h_ = max(ha_mp, 0) * 1000 / area
    rect(s, bx, sy(a), bw, base - sy(a), FLOWER)
    if h_ > 0:
        rect(s, bx, sy(a + h_), bw, sy(a) - sy(a + h_), MOSS)
    text(s, bx - 12, sy(a + h_) - 24, bw + 24, 20, f"{a + h_:.0f}", size=13, bold=True, align=PP_ALIGN.CENTER,
         en=True)
    text(s, bx - 12, base + 6, bw + 24, 20, lab, size=12, color=SECONDARY, align=PP_ALIGN.CENTER, en=True)
rent_y = sy(KAWA_RENT * 1000 / KAWA_TSUBO)
hline(s, x0, rent_y, 1060 - x0, STRONG, 2.0, dash=True)
vline(s, x0 + 10 * slot, top - 10, base - top + 30, BORDER_LIGHT, 1.0, dash=True)
text(s, x0 + 10 * slot - 306, top - 14, 300, 20, "26年4月〜 川崎のみ（新城を閉鎖）→", size=13, color=SECONDARY,
     align=PP_ALIGN.RIGHT)
rect(s, 112, 228, 16, 16, FLOWER)
text(s, 136, 223, 160, 26, "アンドフラワー", size=15, color=SECONDARY)
rect(s, 282, 228, 16, 16, MOSS)
text(s, 306, 223, 120, 26, "ハナイチ", size=15, color=SECONDARY)
hline(s, 412, 236, 28, STRONG, 2.0, dash=True)
text(s, 448, 223, 200, 26, "家賃 7.3千円/坪", size=15, color=SECONDARY)
text(s, 56, 660, 1016, 80, ["平常月は家賃の約4〜5倍、母の日は約17〜19倍。[[ピーク前提の面積]]が平常月に余っている",
                             "→ 平常月の空きを、坪あたり利益の高い用途で埋めることが坪効率向上の鍵"], size=21, bold=True,
     spacing=1.45)
footnote(s, "※ 花拠点_坪単位売上効率_月次推移（25年6月〜26年8月）。25年3月〜26年3月の面積は川崎128.84坪＋新城40.49坪。"
         "限界利益は広告費控除前。観葉（アンドプランツ）はSBSに委託のため対象外")

# ---- 試算の考え方
s = content_slide("坪効率の試算の考え方", "用途別に面積を割り、坪あたり貢献利益で比べる")
steps = [
    ("1", "面積を用途別に割る",
     "図面で専有エリア（AF作業・保管／高付加価値3PL／ハナイチ／単純3PL保管）と共用（梱包・出荷・事務・通路）に分ける。共用は出荷件数比で按分"),
    ("2", "用途別の坪あたり貢献利益を出す",
     "（限界利益 − 直接人件費 − 資材等）÷ 按分後の坪数。月次で出し、家賃の坪単価7,296円と比べる"),
    ("3", "繁忙月と平常月の2断面で見る",
     "母の日など繁忙月に必要な坪数と、平常月に実際に使う坪数の差を「空き坪」として数える"),
    ("4", "空き坪の埋め方を比べて計画にする",
     "単純3PL（保管料＋出荷料）／高付加価値3PL（花同梱・ギフト加工）／AP・AFの拡大を、坪あたり貢献利益の高い順に割り当てる"),
]
y = 168
for n, head, body in steps:
    text(s, 72, y - 6, 70, 70, n, size=52, color=LIGHT, bold=True, en=True)
    text(s, 150, y, 900, 36, head, size=26, bold=True)
    text(s, 150, y + 40, 900, 70, body, size=18, color=SECONDARY, spacing=1.4)
    if n != "4":
        hline(s, 72, y + 118, 1000)
    y += 134
footnote(s, "※ 必要なデータ：拠点の図面（用途別の坪数）、用途別の出荷件数・作業時間（人件費の按分）、3PLの料金表（保管・出荷・加工単価）")

# ---- 用途別の試算（現状の暫定値と計画の枠）
s = content_slide("用途別の試算（たたき台）", "26年8月　単位：千円/月（坪あたりは千円/坪）")
hdr = ["用途", "面積（坪）", "限界利益", "直接人件費", "坪あたり\n貢献利益", "今後の計画"]
cx4 = [56, 300, 440, 580, 720, 880]
wid = [230, 120, 120, 120, 130, 190]
y = 160
for j, c in enumerate(hdr):
    text(s, cx4[j], y, wid[j], 56, c, size=16, color=SECONDARY, bold=True, spacing=1.2,
         align=PP_ALIGN.LEFT if j in (0, 5) else PP_ALIGN.RIGHT)
hline(s, 56, y + 60, 1016, MOSS, 1.5)
y += 72
share_af = 7.186 / (7.186 + 3.285)  # 暫定：8月の売上比で面積を按分
rows = [
    ("アンドフラワー", f"{KAWA_TSUBO * share_af:.0f}*", "4,289", "[　]", f"{4289 / (KAWA_TSUBO * share_af):.0f}*", "平常月の効率を上げる", FLOWER),
    ("ハナイチ", f"{KAWA_TSUBO * (1 - share_af):.0f}*", "450", "[　]", f"{450 / (KAWA_TSUBO * (1 - share_af)):.1f}*", "取扱量の増加", MOSS),
    ("高付加価値3PL\n（花同梱・ギフト）", "[　]", "[　]", "[　]", "[　]", "空き坪に優先配分", MOSS),
    ("単純3PL\n（保管・出荷）", "[　]", "[　]", "[　]", "[　]", "残りの空き坪", MOSS),
    ("共用（梱包・出荷・事務）", "[　]", "—", "—", "—", "出荷件数比で按分", SECONDARY),
]
for name, a, mp, lab, per, plan, col in rows:
    text(s, cx4[0], y, wid[0], 56, name, size=18, bold=True, color=col, spacing=1.2)
    for j, v in enumerate([a, mp, lab, per], 1):
        text(s, cx4[j], y, wid[j], 30, v, size=19, align=PP_ALIGN.RIGHT, en=True,
             color=MUTED if v.startswith("[") else MOSS, bold=(j == 4))
    text(s, cx4[5], y, wid[5], 56, plan, size=16, color=SECONDARY, spacing=1.2)
    y += 70
    hline(s, 56, y - 10, 1016)
text(s, 56, y + 6, 1016, 80,
     [f"暫定でも、アンドフラワーとハナイチで[[坪あたり利益に約{(4289 / share_af) / (450 / (1 - share_af)):.0f}倍の差]]がある",
      "図面で実際の面積を割り、3PLの単価を入れると、空き坪の使い道が決められる"], size=21, bold=True, spacing=1.45)
footnote(s, "* 面積は暫定で8月の売上比（AF 69%／ハナイチ 31%）で按分。[　] は今後入れる数値。家賃の坪単価は7.3千円/坪")

# ---- 次のステップ
s = content_slide("今後の進め方")
nexts = [
    ("ハナイチの直接固定費を確定する", "人件費（現場・CS）と設備費を、AP/AFと同じ貢献利益の形にそろえる"),
    ("川崎拠点の用途別の坪数を測る", "図面で専有と共用に分け、出荷件数で共用を按分する"),
    ("3PLの単価をそろえる", "単純3PL（保管・出荷）と高付加価値3PL（花同梱・加工）の料金表から坪あたり利益を出す"),
    ("Value UPプランを数値に落とす", "卸・輸入で下がる原価率と、売価に回す幅を事業ごとに置く"),
]
y = 170
for i, (h_, b_) in enumerate(nexts, 1):
    text(s, 96, y, 60, 60, str(i), size=44, color=LIGHT, bold=True, en=True)
    text(s, 170, y + 2, 880, 36, h_, size=28, bold=True)
    text(s, 170, y + 46, 880, 30, b_, size=19, color=SECONDARY)
    if i < len(nexts):
        hline(s, 96, y + 104, 936)
    y += 128

# ---- 要約（2ページ目の次。数字が出そろってから描く）
s = summary_slide
SUM = [
    ("1", "各事業のエコノミクスの推移", [
        f"アンドプランツ：26年9月期の売上{ap_fy_s[3]:.0f}百万円は前期並み。広告費を{ap_fy_ad[2]:.0f} → {ap_fy_ad[3]:.0f}百万円に抑え、"
        f"広告費込み限界利益は{ap_fy_net[3]:.0f}百万円（+{(ap_fy_net[3] / ap_fy_net[2] - 1) * 100:.0f}%）",
        f"アンドフラワー：売上{af_fy_s[3]:.0f}百万円。母の日に集中し、広告費込み限界利益は{af_fy_net[2]:.0f} → {af_fy_net[3]:.0f}百万円",
        f"ハナイチ：売上{ha_fy_s[3]:.0f}百万円（+{(ha_fy_s[3] / ha_fy_s[2] - 1) * 100:.0f}%）。注文数が伸びている"]),
    ("2", "今後のエコノミクス", [
        f"貢献利益（月）は、アンドプランツが閑散期{PERIODS['sep']['ap']['CONTRIB']:.1f}・繁忙期{PERIODS['may']['ap']['CONTRIB']:.1f}百万円。"
        f"アンドフラワーは閑散期▲{-PERIODS['sep']['af']['CONTRIB']:.1f}、繁忙期▲{-PERIODS['may']['af']['CONTRIB']:.1f}百万円",
        f"アンドフラワーの損益分岐は月商 約{af_be:.0f}百万円。卸・輸入で原価率を5pt下げると年+{tot5:.0f}百万円（3事業計）"]),
    ("3", "倉庫の収益性", [
        "川崎拠点の坪あたり限界利益は、平常月で家賃の約4〜5倍、母の日の月は約17〜19倍",
        "ピークに合わせた面積が平常月に余っている。用途別に面積を割り、空き坪を坪あたり利益の高い用途で埋める"]),
]
y = 146
for n, head, lines_ in SUM:
    text(s, 56, y, 60, 50, n, size=44, color=LIGHT, bold=True, en=True)
    text(s, 120, y + 6, 940, 36, head, size=27, bold=True)
    y += 52
    for ln_ in lines_:
        rect(s, 124, y + 16, 8, 8, LIGHT)  # 行間の余白は文字の上に付くので、1行目の文字の中心に合わせて下げる
        text(s, 142, y, 920, 64, ln_, size=19, spacing=1.45)
        y += 62 if len(ln_) > 46 else 36
    y += 26
footnote(s, "※ 数値は税抜。用語の定義と集計の前提は次のページ")

prs.save("giftee_economics_2026-10.pptx")  # PDF は soffice で書き出す
print("saved", len(prs.slides._sldIdLst), "slides")
for name, fs, fa, fn in [("AP", ap_fy_s, ap_fy_ad, ap_fy_net), ("AF", af_fy_s, af_fy_ad, af_fy_net),
                         ("HA", ha_fy_s, ha_fy_ad, ha_fy_net)]:
    print(name, [round(v, 1) for v in fs], [round(v, 1) for v in fa], [round(v, 1) for v in fn])
print({p_: {b_: {kk: round(vv, 2) for kk, vv in PERIODS[p_][b_].items()} for b_ in ("ap", "af", "ha")} for p_ in PERIODS})
